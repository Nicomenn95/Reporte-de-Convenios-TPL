"""TPL | Liquidación de convenios. Ejecutar: streamlit run app.py

V1: liquidación por fecha de venta y estado al descargar el reporte.
No reconstruye movimientos históricos ni saldos de facturas anteriores.
"""
from __future__ import annotations

import hashlib
import io
import re
import unicodedata
import zipfile
from datetime import datetime
from xml.sax.saxutils import escape

import numpy as np
import pandas as pd
import plotly.express as px
import streamlit as st
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import (
    LongTable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
)

REQUIRED = [
    "Agencia", "Estación de trabajo", "Fecha y hora de venta", "Fecha salida",
    "Hora salida", "Código pasaje", "Valor de venta", "Moneda", "Estado",
    "Origen (ciudad)", "Destino (ciudad)",
]
DETAIL = [
    "Empresa", "Centro de costo", "Fecha venta", "Código pasaje",
    "Código autorización", "Estado", "Ticket en blanco", "Origen", "Destino",
    "Fecha viaje", "Hora viaje", "Pasajero", "ID pasajero", "Usuario",
    "Clase asiento", "Asiento", "Valor original", "Monto excluido",
    "Monto a cobrar", "Criterio",
]
NOTE = (
    "Liquidación por fecha de venta y estado informado al descargar el reporte. "
    "Pagado y Cambiado conservan el valor original; Devuelto, Anulado y códigos "
    "terminados en TB tienen monto a cobrar $0. Los cambios no generan un segundo cargo. "
    "No incluye saldos anteriores, pagos de facturas ni ajustes entre meses. "
    "El estado Pagado del pasaje no acredita el pago de la factura del cliente."
)


def normalized(value):
    return " ".join(str(value).strip().split())


def key(value):
    text = unicodedata.normalize("NFKD", normalized(value))
    return "".join(c for c in text if not unicodedata.combining(c)).casefold()


def money(value):
    return "$" + f"{int(value):,}".replace(",", ".")


def filename(value):
    value = unicodedata.normalize("NFKD", str(value)).encode("ascii", "ignore").decode()
    return re.sub(r"[^A-Za-z0-9_-]+", "_", value).strip("_")[:85] or "empresa"


def parse_amount(value):
    if pd.isna(value):
        return np.nan
    if isinstance(value, (int, float, np.number)):
        return float(value)
    text = re.sub(r"[\s$]", "", str(value)).replace("CLP", "")
    # Miles chilenos: 8.500 o 1.234.500; decimales: 8500,00.
    if re.fullmatch(r"[+-]?\d{1,3}(\.\d{3})+(,\d+)?", text):
        text = text.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"[+-]?\d{1,3}(,\d{3})+", text):
        text = text.replace(",", "")
    else:
        text = text.replace(",", ".")
    return pd.to_numeric(text, errors="coerce")


def read_source(data: bytes, name: str):
    """Detecta hoja y encabezado; no depende de un nombre fijo de hoja."""
    if len(data) > 30 * 1024 * 1024:
        raise ValueError("El archivo supera el límite de 30 MB de esta aplicación.")
    if name.lower().endswith(".csv"):
        raw = None
        for encoding in ("utf-8-sig", "cp1252"):
            try:
                raw = pd.read_csv(io.BytesIO(data), header=None, sep=None,
                                  engine="python", encoding=encoding, dtype=object)
                break
            except UnicodeDecodeError:
                continue
        if raw is None:
            raise ValueError("No fue posible leer la codificación del CSV.")
        sheets = [("CSV", raw)]
    else:
        book = pd.ExcelFile(io.BytesIO(data))
        sheets = [(sheet, book.parse(sheet, header=None, dtype=object))
                  for sheet in book.sheet_names]
    required = {key(c) for c in REQUIRED}
    best = set()
    for sheet, raw in sheets:
        for pos in range(min(15, len(raw))):
            headers = [normalized(c) if pd.notna(c) else "" for c in raw.iloc[pos]]
            found = {key(c) for c in headers}
            if len(required & found) > len(best):
                best = required & found
            if required <= found:
                canon = {key(c): c for c in REQUIRED + [
                    "Código autorización", "Primer nombre", "Apellido paterno",
                    "ID del pasajero", "Usuario", "Clase asiento", "Asiento"]}
                headers = [canon.get(key(c), c or f"Columna_{i}")
                           for i, c in enumerate(headers)]
                if len(headers) != len(set(headers)):
                    raise ValueError("Hay encabezados duplicados en la hoja de convenios.")
                df = raw.iloc[pos + 1:].copy()
                df.columns = headers
                df = df.dropna(how="all").reset_index(drop=True)
                return df, sheet
    missing = [c for c in REQUIRED if key(c) not in best]
    raise ValueError("No se encontró una tabla de convenios válida. Faltan: " + ", ".join(missing))


def prepare(raw):
    df = raw.copy()
    text_cols = list(set(REQUIRED) - {"Valor de venta", "Fecha y hora de venta", "Fecha salida"})
    for col in text_cols:
        df[col] = df[col].fillna("").map(normalized)
    for col in ["Agencia", "Código pasaje"]:
        if df[col].eq("").any():
            raise ValueError(f"Hay filas sin {col}. Corrige el archivo antes de generar.")
    df["Estación de trabajo"] = df["Estación de trabajo"].replace("", "SIN CENTRO DE COSTO")
    if df.empty:
        raise ValueError("El reporte no contiene registros.")
    duplicates = df["Código pasaje"].str.upper().duplicated(keep=False)
    if duplicates.any():
        codes = df.loc[duplicates, "Código pasaje"].unique()[:8]
        raise ValueError("Hay códigos de pasaje repetidos. No se sumarán automáticamente: " + ", ".join(codes))
    states = {"pagado": "Pagado", "cambiado": "Cambiado", "devuelto": "Devuelto", "anulado": "Anulado"}
    state = df["Estado"].map(key).map(states)
    if state.isna().any():
        raise ValueError("Estados no reconocidos: " + ", ".join(df.loc[state.isna(), "Estado"].unique()))
    df["Estado"] = state
    if not df["Moneda"].str.upper().eq("CLP").all():
        raise ValueError("Esta V1 solo procesa CLP; el archivo contiene otra moneda o moneda vacía.")
    amount = df["Valor de venta"].map(parse_amount)
    invalid = ~np.isfinite(amount) | amount.lt(0) | amount.gt(1e12) | amount.mod(1).ne(0)
    if invalid.any():
        raise ValueError("Hay valores de venta vacíos, negativos, no numéricos o con decimales. Se esperan pesos CLP enteros.")
    sale = pd.to_datetime(raw["Fecha y hora de venta"], format="mixed", dayfirst=True, errors="coerce")
    travel = pd.to_datetime(raw["Fecha salida"], format="mixed", dayfirst=True, errors="coerce")
    if sale.isna().any() or travel.isna().any():
        raise ValueError("Hay fechas de venta o salida inválidas. Usa día/mes/año o fechas Excel.")
    out = pd.DataFrame(index=df.index)
    out["Empresa"] = df["Agencia"]
    out["Centro de costo"] = df["Estación de trabajo"]
    out["Fecha venta"] = sale
    out["Código pasaje"] = df["Código pasaje"]
    out["Estado"] = state
    tb = df["Código pasaje"].str.upper().str.endswith("TB")
    out["Ticket en blanco"] = np.where(tb, "Sí", "No")
    out["Origen"] = df["Origen (ciudad)"]
    out["Destino"] = df["Destino (ciudad)"]
    out["Fecha viaje"] = travel
    out["Hora viaje"] = df["Hora salida"]
    for target, source in {"Código autorización": "Código autorización", "ID pasajero": "ID del pasajero",
                           "Usuario": "Usuario", "Clase asiento": "Clase asiento", "Asiento": "Asiento"}.items():
        out[target] = df.get(source, pd.Series("", index=df.index)).astype("string").fillna("").astype(str)
    first = df.get("Primer nombre", pd.Series("", index=df.index)).fillna("").astype(str)
    last = df.get("Apellido paterno", pd.Series("", index=df.index)).fillna("").astype(str)
    out["Pasajero"] = (first + " " + last).map(normalized)
    out["Valor original"] = amount.astype("int64")
    eligible = ~tb & state.isin(["Pagado", "Cambiado"])
    out["Monto a cobrar"] = np.where(eligible, out["Valor original"], 0).astype("int64")
    out["Monto excluido"] = out["Valor original"] - out["Monto a cobrar"]
    out["Criterio"] = np.select(
        [tb, state.eq("Devuelto"), state.eq("Anulado"), state.eq("Cambiado")],
        ["Ticket en blanco excluido", "Devuelto: excluido", "Anulado: excluido",
         "Cambio: se conserva la compra original"], default="Compra incluida")
    return out[DETAIL].sort_values(["Empresa", "Centro de costo", "Fecha venta", "Código pasaje"]).reset_index(drop=True)


def summary(df, groups):
    table = df.groupby(groups, dropna=False, sort=True).agg(
        Pasajes=("Código pasaje", "size"),
        **{"Valor original": ("Valor original", "sum"), "Monto excluido": ("Monto excluido", "sum"),
           "Monto a cobrar": ("Monto a cobrar", "sum")}).reset_index()
    return table


def period_text(df):
    return f"{df['Fecha venta'].min():%d/%m/%Y} a {df['Fecha venta'].max():%d/%m/%Y} (ventas incluidas)"


def excel_report(df, period, source):
    buffer = io.BytesIO()
    tables = {
        "Resumen": summary(df, ["Empresa", "Centro de costo"]),
        "Detalle": df.copy(),
        "Estados": summary(df, ["Empresa", "Estado", "Ticket en blanco"]),
        "Criterios": pd.DataFrame({"Campo": ["Período de venta", "Fuente", "Alcance", "Registros", "Total a cobrar"],
                                   "Contenido": [period, source, NOTE, len(df), int(df["Monto a cobrar"].sum())]}),
    }
    with pd.ExcelWriter(buffer, engine="openpyxl", datetime_format="DD/MM/YYYY HH:MM", date_format="DD/MM/YYYY") as writer:
        for name, table in tables.items():
            table = table.copy()
            # Evita que texto proveniente del archivo se convierta en fórmulas.
            for col in table.select_dtypes(include=["object", "string"]):
                table[col] = table[col].map(lambda v: "'" + v if isinstance(v, str) and v.startswith(("=", "+", "-", "@")) else v)
            table.to_excel(writer, sheet_name=name, index=False)
            ws = writer.sheets[name]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="16324F")
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.row_dimensions[1].height = 32
            for j, col in enumerate(table.columns, 1):
                samples = table[col].head(200).astype(str).map(len)
                width = min(46, max(14, len(str(col)) + 2, int(samples.max()) + 2 if len(samples) else 14))
                ws.column_dimensions[get_column_letter(j)].width = width
                for row in range(2, ws.max_row + 1):
                    cell = ws.cell(row, j)
                    if row % 2 == 0:
                        cell.fill = PatternFill("solid", fgColor="F0F4F8")
                    if col in ("Valor original", "Monto excluido", "Monto a cobrar"):
                        cell.number_format = '"$"#,##0'
                    if col == "Fecha viaje":
                        cell.number_format = "DD/MM/YYYY"
            if name in ("Resumen", "Estados"):
                total_row = ws.max_row + 1
                ws.cell(total_row, 1, "TOTAL")
                for j, col in enumerate(table.columns, 1):
                    if col in ("Pasajes", "Valor original", "Monto excluido", "Monto a cobrar"):
                        cell = ws.cell(total_row, j, int(table[col].sum()))
                        cell.number_format = '#,##0' if col == "Pasajes" else '"$"#,##0'
                for cell in ws[total_row]:
                    cell.font = Font(bold=True)
            if name == "Criterios":
                ws.column_dimensions["B"].width = 110
                ws["B4"].alignment = Alignment(wrap_text=True, vertical="top")
                ws.row_dimensions[4].height = 75
    return buffer.getvalue()


def pdf_report(df, company, period, source):
    """Cartola de liquidación: muestra estados, no inventa eventos históricos."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), rightMargin=25, leftMargin=25,
                            topMargin=30, bottomMargin=35, title=f"Liquidación {company}")
    styles = getSampleStyleSheet()
    small = ParagraphStyle("Cell", fontName="Helvetica", fontSize=6.5, leading=8)
    head = ParagraphStyle("Head", parent=small, textColor=colors.white, fontName="Helvetica-Bold")
    right = ParagraphStyle("Right", parent=small, alignment=TA_RIGHT)
    def p(value, style=small):
        return Paragraph(escape(str(value)), style)
    story = [p("TPL | Cartola de liquidación de convenios", styles["Title"]),
             p(company, styles["Heading2"]), p(f"Período de venta: {period}"),
             p(f"Archivo fuente: {source}"), Spacer(1, 8), p(NOTE), Spacer(1, 10)]
    sums = summary(df, ["Centro de costo"])
    data = [[p(c, head) for c in ["Centro de costo", "Pasajes", "Valor original", "Excluido", "A cobrar"]]]
    for _, row in sums.iterrows():
        data.append([p(row["Centro de costo"]), p(row["Pasajes"])] +
                    [p(money(row[c]), right) for c in ["Valor original", "Monto excluido", "Monto a cobrar"]])
    summary_weights = [300, 55, 140, 140, 150]
    table = Table(data, colWidths=[doc.width * z / sum(summary_weights) for z in summary_weights], repeatRows=1, hAlign="LEFT")
    base_style = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#16324F")),
                  ("VALIGN", (0, 0), (-1, -1), "TOP"),
                  ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F0F4F8")]),
                  ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]
    table.setStyle(TableStyle(base_style))
    story.extend([table, Spacer(1, 10), p("Total a cobrar: " + money(df["Monto a cobrar"].sum()), styles["Heading2"])])
    for center, group in df.groupby("Centro de costo", sort=True):
        story.extend([Spacer(1, 8), p("Centro de costo: " + center, styles["Heading2"])])
        labels = ["Venta", "Código pasaje", "Estado / criterio", "Origen", "Destino", "Viaje", "Pasajero", "Original", "A cobrar"]
        rows = [[p(c, head) for c in labels]]
        for _, r in group.iterrows():
            status = r["Estado"] + (" / TB excluido" if r["Ticket en blanco"] == "Sí" else "")
            values = [r["Fecha venta"].strftime("%d/%m/%y %H:%M"), r["Código pasaje"], status,
                      r["Origen"], r["Destino"], r["Fecha viaje"].strftime("%d/%m/%y") + " " + r["Hora viaje"],
                      r["Pasajero"], money(r["Valor original"]), money(r["Monto a cobrar"])]
            rows.append([p(value, right if j >= 7 else small) for j, value in enumerate(values)])
        weights = [69, 105, 68, 61, 61, 80, 109, 65, 65]
        t = LongTable(rows, colWidths=[doc.width * z / sum(weights) for z in weights], repeatRows=1, hAlign="LEFT")
        t.setStyle(TableStyle(base_style))
        story.extend([t, Spacer(1, 5), p("Subtotal centro de costo: " + money(group["Monto a cobrar"].sum()), styles["Heading3"])])
    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.drawString(25, 18, "TPL | Liquidación de convenios | " + period)
        canvas.drawRightString(landscape(A4)[0] - 25, 18, f"Página {document.page}")
        canvas.restoreState()
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()


def company_stem(company, period):
    # Hash evita colisiones entre razones sociales con el mismo nombre saneado.
    suffix = hashlib.sha256(company.encode()).hexdigest()[:6]
    return f"Convenios_{filename(company)}_{suffix}_{period}"


def main():
    st.set_page_config(page_title="TPL | Convenios", page_icon="🚌", layout="wide")
    st.title("Convenios TPL")
    st.caption("Liquidación por cliente y centro de costo · V1")
    with st.expander("Reglas de cobro y alcance"):
        st.write(NOTE)
        st.write("El sufijo TB se excluye siempre, aunque el estado diga Pagado. No se requiere base de datos para esta liquidación por archivo.")
    upload = st.file_uploader("Carga el reporte original de convenios de Terra", type=["xlsx", "xls", "csv"])
    if upload is None:
        st.info("Carga un archivo para revisar el cálculo y descargar las liquidaciones.")
        return
    data = upload.getvalue()
    token = hashlib.sha256(data).hexdigest()
    if st.session_state.get("source_token") != token:
        st.session_state.pop("bundle", None)
        try:
            raw, sheet = read_source(data, upload.name)
            prepared = prepare(raw)
        except Exception as exc:
            st.error(f"No se pudo procesar el reporte: {exc}")
            return
        st.session_state.update(source_token=token, prepared=prepared, source_sheet=sheet)
    full = st.session_state.prepared
    months = sorted(full["Fecha venta"].dt.strftime("%Y-%m").unique(), reverse=True)
    with st.sidebar:
        st.header("Período y cliente")
        month = st.selectbox("Mes de venta", months)
        monthly = full.loc[full["Fecha venta"].dt.strftime("%Y-%m").eq(month)]
        companies = sorted(monthly["Empresa"].unique())
        company = st.selectbox("Cliente empresa", companies)
        st.caption("Las descargas individuales incluyen todos los centros de costo del cliente seleccionado.")
    selected = monthly.loc[monthly["Empresa"].eq(company)].copy()
    st.caption(f"Fuente: {upload.name} · Hoja: {st.session_state.source_sheet} · {len(monthly):,} pasajes en {month}")
    cards = st.columns(4)
    cards[0].metric("A cobrar · mes completo", money(monthly["Monto a cobrar"].sum()))
    cards[1].metric("Clientes", monthly["Empresa"].nunique())
    cards[2].metric("Pasajes", len(monthly))
    cards[3].metric("Valor excluido", money(monthly["Monto excluido"].sum()))
    overview, client_tab, downloads, review = st.tabs(["Resumen mensual", "Cliente y detalle", "Descargas", "Validaciones"])
    with overview:
        totals = summary(monthly, ["Empresa"])
        st.dataframe(totals.style.map(lambda v: "font-weight: bold; color: #166534" if v > 0 else "color: #64748b", subset=["Monto a cobrar"]).format({c: money for c in ["Valor original", "Monto excluido", "Monto a cobrar"]}), hide_index=True, width="stretch")
        fig = px.bar(totals, x="Monto a cobrar", y="Empresa", orientation="h", title="Monto a cobrar por cliente", color_discrete_sequence=["#176B87"])
        fig.update_layout(yaxis={"categoryorder": "total ascending"}, xaxis_title="Pesos CLP")
        st.plotly_chart(fig, width="stretch")
    with client_tab:
        st.subheader(company)
        st.metric("Total a cobrar al cliente", money(selected["Monto a cobrar"].sum()))
        st.dataframe(summary(selected, ["Centro de costo"]), hide_index=True, width="stretch")
        st.caption("Detalle completo del cliente. Los importes son CLP.")
        st.dataframe(selected, hide_index=True, width="stretch")
    with downloads:
        st.subheader("Archivos del cliente seleccionado")
        if st.button("Preparar Excel y PDF del cliente", type="primary"):
            with st.spinner("Generando archivos..."):
                try:
                    xlsx = excel_report(selected, month, upload.name)
                    pdf = pdf_report(selected, company, month, upload.name)
                    st.session_state["bundle"] = (token, month, company, xlsx, pdf)
                except Exception as exc:
                    st.error(f"No se pudieron generar los archivos: {exc}")
        bundle = st.session_state.get("bundle")
        if bundle and bundle[:3] == (token, month, company):
            stem = company_stem(company, month)
            left, right_col = st.columns(2)
            left.download_button("Descargar Excel", bundle[3], stem + ".xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", on_click="ignore", width="stretch")
            right_col.download_button("Descargar cartola PDF", bundle[4], stem + ".pdf", mime="application/pdf", on_click="ignore", width="stretch")
        st.divider()
        st.subheader("Todas las empresas del mes")
        zip_key = (token, month)
        if st.button("Preparar ZIP con todos los Excel y PDF"):
            with st.spinner("Generando liquidaciones de todas las empresas..."):
                try:
                    buffer = io.BytesIO()
                    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
                        archive.writestr("Resumen_" + month + ".xlsx", excel_report(monthly, month, upload.name))
                        for name, part in monthly.groupby("Empresa", sort=True):
                            stem = company_stem(name, month)
                            archive.writestr(stem + ".xlsx", excel_report(part, month, upload.name))
                            archive.writestr(stem + ".pdf", pdf_report(part, name, month, upload.name))
                    st.session_state["zip_bundle"] = (zip_key, buffer.getvalue())
                except Exception as exc:
                    st.error(f"No se pudo completar el ZIP: {exc}")
        zipped = st.session_state.get("zip_bundle")
        if zipped and zipped[0] == zip_key:
            st.download_button("Descargar todas las liquidaciones", zipped[1], f"Convenios_TPL_{month}.zip", mime="application/zip", on_click="ignore", width="stretch")
    with review:
        st.success("Columnas, códigos únicos, estados, fechas e importes validados.")
        st.write(f"Fechas presentes: {period_text(monthly)}. La carga no acredita que se hayan exportado todos los días del mes.")
        st.info("El reporte no contiene la fecha de devolución. Esta V1 excluye los devueltos/anulados según su estado actual, aunque su devolución haya ocurrido después del mes de venta.")
        tb = monthly.loc[monthly["Ticket en blanco"].eq("Sí")]
        st.write(f"Tickets en blanco excluidos: {len(tb)} · valor original {money(tb['Valor original'].sum())}")
        st.dataframe(summary(monthly, ["Estado", "Ticket en blanco"]), hide_index=True, width="stretch")
        if not tb.empty:
            st.dataframe(tb[["Empresa", "Código pasaje", "Valor original", "Monto a cobrar", "Criterio"]], hide_index=True, width="stretch")
        st.caption("Los archivos cargados se procesan en memoria de la sesión. No se escriben datos de pasajeros en el repositorio.")


if __name__ == "__main__":
    main()
